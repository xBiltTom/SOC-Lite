"""Consulta, verificación y exportación de la bitácora de auditoría."""
from __future__ import annotations

import csv
import io
import secrets
from datetime import datetime
from typing import Annotated, Any

from fastapi import APIRouter, Depends, Header, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select, text

from app.core.config import settings
from app.core.deps import DbSession, PaginationDep, require_permission
from app.core.permissions import P
from app.models.audit import AuditLog
from app.models.enums import AuditResult
from app.schemas.common import Page
from app.schemas.report import AuditEventIn, AuditOut, AuditStats, ChainVerify
from app.services import audit

router = APIRouter(prefix="/audit", tags=["Auditoría"])

ReadAudit = Annotated[Any, Depends(require_permission(P.AUDIT_READ))]


def _filters(
    module: str | None,
    action: str | None,
    username: str | None,
    result: AuditResult | None,
    date_from: datetime | None,
    date_to: datetime | None,
    q: str | None,
):
    conditions = []
    if module:
        conditions.append(AuditLog.module == module)
    if action:
        conditions.append(AuditLog.action == action)
    if username:
        conditions.append(AuditLog.username.ilike(f"%{username}%"))
    if result:
        conditions.append(AuditLog.result == result)
    if date_from:
        conditions.append(AuditLog.occurred_at >= date_from)
    if date_to:
        conditions.append(AuditLog.occurred_at <= date_to)
    if q:
        pattern = f"%{q}%"
        conditions.append(
            or_(
                AuditLog.entity_id.ilike(pattern),
                AuditLog.request_id.ilike(pattern),
                AuditLog.username.ilike(pattern),
                text("CAST(audit_log.detail AS text) ILIKE :q").bindparams(q=pattern),
            )
        )
    return conditions


@router.get("", response_model=Page[AuditOut], summary="Consultar bitácora")
async def list_audit(
    db: DbSession,
    _: ReadAudit,
    pagination: PaginationDep,
    module: str | None = Query(default=None, max_length=40),
    action: str | None = Query(default=None, max_length=60),
    username: str | None = Query(default=None, max_length=60),
    result: AuditResult | None = None,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    q: str | None = Query(default=None, max_length=120),
) -> Page[AuditOut]:
    conditions = _filters(module, action, username, result, date_from, date_to, q)
    total = (
        await db.execute(select(func.count(AuditLog.id)).where(*conditions))
    ).scalar_one()
    rows = (
        await db.execute(
            select(AuditLog)
            .where(*conditions)
            .order_by(AuditLog.chain_seq.desc())
            .offset(pagination.offset)
            .limit(pagination.size)
        )
    ).scalars().all()
    return Page[AuditOut](
        items=[AuditOut.model_validate(row) for row in rows],
        total=total,
        page=pagination.page,
        size=pagination.size,
        pages=max(1, -(-total // pagination.size)),
    )


@router.post("/events", status_code=201, summary="Registrar evento desde automatización (n8n)")
async def ingest_event(
    payload: AuditEventIn,
    request: Request,
    db: DbSession,
    x_ingest_key: str | None = Header(default=None, alias="X-Ingest-Key"),
) -> dict[str, Any]:
    if not x_ingest_key or not secrets.compare_digest(x_ingest_key, settings.ingest_api_key):
        raise Unauthorized("Clave de ingesta inválida")
    entry = await audit.record(
        db,
        module=payload.module,
        action=payload.action,
        username=payload.username or "n8n",
        request=request,
        entity_type=payload.entity_type,
        entity_id=payload.entity_id,
        result=payload.result,
        detail=payload.detail,
    )
    await db.commit()
    return {"id": entry.id, "hash": entry.hash}


@router.get("/modules", response_model=list[dict[str, Any]], summary="Módulos auditados")
async def modules(db: DbSession, _: ReadAudit) -> list[dict[str, Any]]:
    rows = (
        await db.execute(
            text(
                "SELECT module, COUNT(*) AS total FROM audit_log GROUP BY module ORDER BY total DESC"
            )
        )
    ).mappings().all()
    return [{"module": row["module"], "total": int(row["total"])} for row in rows]


@router.get("/stats", response_model=AuditStats, summary="Resumen de la bitácora")
async def stats(db: DbSession, _: ReadAudit, days: int = Query(default=30, ge=1, le=365)) -> AuditStats:
    row = (
        await db.execute(
            text(
                """
                SELECT COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE result = 'success') AS success,
                       COUNT(*) FILTER (WHERE result = 'failure') AS failure,
                       COUNT(*) FILTER (WHERE result = 'denied') AS denied
                FROM audit_log WHERE occurred_at >= now() - make_interval(days => :days)
                """
            ),
            {"days": days},
        )
    ).mappings().one()
    by_module = (
        await db.execute(
            text(
                """
                SELECT module, COUNT(*) AS total FROM audit_log
                WHERE occurred_at >= now() - make_interval(days => :days)
                GROUP BY module ORDER BY total DESC
                """
            ),
            {"days": days},
        )
    ).mappings().all()
    return AuditStats(
        total=int(row["total"]),
        success=int(row["success"]),
        failure=int(row["failure"]),
        denied=int(row["denied"]),
        modules=[{"module": m["module"], "total": int(m["total"])} for m in by_module],
    )


@router.get("/verify-chain", response_model=ChainVerify, summary="Verificar integridad de la cadena de hashes")
async def verify_chain(db: DbSession, _: ReadAudit) -> ChainVerify:
    valid, checked, broken_at = await audit.verify_chain(db)
    return ChainVerify(
        valid=valid,
        checked=checked,
        broken_at=broken_at,
        message="Cadena íntegra" if valid else "Se detectó una alteración en la bitácora",
    )


@router.get("/export", summary="Exportar bitácora (CSV/XLSX)")
async def export_audit(
    request: Request,
    db: DbSession,
    actor: Annotated[Any, Depends(require_permission(P.AUDIT_EXPORT))],
    fmt: str = Query(default="csv", alias="format", pattern="^(csv|xlsx)$"),
    module: str | None = None,
    action: str | None = None,
    username: str | None = None,
    result: AuditResult | None = None,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    q: str | None = None,
    limit: int = Query(default=10000, ge=1, le=50000),
) -> StreamingResponse:
    conditions = _filters(module, action, username, result, date_from, date_to, q)
    rows = (
        await db.execute(
            select(AuditLog).where(*conditions).order_by(AuditLog.chain_seq.desc()).limit(limit)
        )
    ).scalars().all()

    await audit.record(
        db, module="audit", action="AUDIT_EXPORT", user=actor, request=request,
        entity_type="audit_log", detail={"format": fmt, "rows": len(rows), "limit": limit},
    )
    await db.commit()

    headers = ["id", "occurred_at", "username", "module", "action", "entity_type",
               "entity_id", "result", "ip", "request_id", "duration_ms", "hash"]
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(
                [
                    row.id,
                    row.occurred_at.isoformat(),
                    row.username or "",
                    row.module,
                    row.action,
                    row.entity_type or "",
                    row.entity_id or "",
                    row.result.value,
                    str(row.ip) if row.ip else "",
                    row.request_id or "",
                    row.duration_ms if row.duration_ms is not None else "",
                    row.hash or "",
                ]
            )
        payload = io.BytesIO(buffer.getvalue().encode("utf-8-sig"))
        media = "text/csv"
        filename = "soc-lite-auditoria.csv"
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Auditoría"
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
        for row in rows:
            ws.append(
                [
                    row.id,
                    row.occurred_at.strftime("%Y-%m-%d %H:%M:%S"),
                    row.username or "",
                    row.module,
                    row.action,
                    row.entity_type or "",
                    row.entity_id or "",
                    row.result.value,
                    str(row.ip) if row.ip else "",
                    row.request_id or "",
                    row.duration_ms if row.duration_ms is not None else "",
                    row.hash or "",
                ]
            )
        ws.freeze_panes = "A2"
        widths = [8, 20, 16, 14, 22, 14, 24, 12, 16, 20, 12, 68]
        from openpyxl.utils import get_column_letter

        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        payload = io.BytesIO()
        wb.save(payload)
        payload.seek(0)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = "soc-lite-auditoria.xlsx"

    payload.seek(0)
    return StreamingResponse(
        payload,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
