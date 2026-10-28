#!/usr/bin/env python3
"""Construye DOCUMENTACION_SOC-Lite.docx a partir de los Markdown del proyecto.

- Renderiza diagramas Mermaid con mermaid.ink (si hay red) y los inserta como imagen.
- Convierte tablas, listas, bloques de código, citas e imágenes.
- Añade portada, ficha de cumplimiento, encabezado y numeración de páginas.
"""
from __future__ import annotations

import base64
import hashlib
import re
import sys
import urllib.request
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
DIAGRAMS.mkdir(exist_ok=True)
OUT = ROOT / "DOCUMENTACION_SOC-Lite.docx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREY = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0x64, 0x74, 0x8B)
CODE_BG = "F1F5F9"
HEADER_BG = "1E3A8A"


# ---------------------------------------------------------------------------
# Utilidades de formato
# ---------------------------------------------------------------------------
def shade(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def cell_shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1D4ED8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, size: float | None = None, color: RGBColor | None = None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()  # type: ignore[union-attr]
            add_hyperlink(paragraph, label, url)
            pos = match.end()
            continue
        if size:
            run.font.size = Pt(size)
        if color and not token.startswith("["):
            run.font.color.rgb = color
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def render_mermaid(code: str) -> Path | None:
    digest = hashlib.sha256(code.encode("utf-8")).hexdigest()[:16]
    cached = DIAGRAMS / f"mermaid-{digest}.jpg"
    if cached.exists() and cached.stat().st_size > 500:
        return cached
    encoded = base64.urlsafe_b64encode(code.encode("utf-8")).decode()
    url = f"https://mermaid.ink/img/{encoded}?type=png&bgColor=FFFFFF"
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 SOC-Lite-doc"})
    try:
        with urllib.request.urlopen(request, timeout=40) as response:
            payload = response.read()
        if len(payload) < 500:
            return None
        cached.write_bytes(payload)
        return cached
    except Exception as exc:  # sin red o diagrama inválido: se deja como código
        print(f"  [mermaid] no se pudo renderizar ({exc}); se inserta como código")
        return None


def add_code_block(doc: Document, code: str) -> None:
    for chunk in code.rstrip().split("\n"):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.left_indent = Cm(0.25)
        run = paragraph.add_run(chunk if chunk else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        shade(paragraph, CODE_BG)


def add_image(doc: Document, image_path: Path, caption: str | None = None, width: float = 6.2) -> None:
    if not image_path.exists():
        doc.add_paragraph(f"[imagen no encontrada: {image_path.name}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT


def split_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    line = line.replace("\\|", "§PIPE§")
    cells = [cell.strip().replace("§PIPE§", "|") for cell in line.split("|")]
    return cells


def is_separator(line: str) -> bool:
    return bool(re.match(r"^\s*\|?[\s:|-]+\|?\s*$", line)) and "-" in line and "|" in line


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    header_cells = table.rows[0].cells
    for index, text in enumerate(header):
        cell = header_cells[index]
        cell_shade(cell, HEADER_BG)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text.replace("**", "").replace("`", ""))
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row_index, row in enumerate(rows):
        row_obj = table.add_row()
        row_tr_pr = row_obj._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_tr_pr.append(cant_split)
        cells = row_obj.cells
        for index, text in enumerate(row[: len(header)]):
            cell = cells[index]
            set_cell_margins(cell)
            if row_index % 2 == 1:
                cell_shade(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, text, size=8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------------------
# Parser Markdown
# ---------------------------------------------------------------------------
def parse_markdown(doc: Document, path: Path, figure_start: int = 1) -> int:
    lines = path.read_text(encoding="utf-8").splitlines()
    base = path.parent
    index = 0
    figure = figure_start
    in_code = False
    code_lines: list[str] = []
    code_lang = ""
    first_heading = True

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # Bloques de código
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip().lower()
                code_lines = []
            else:
                in_code = False
                code = "\n".join(code_lines)
                if code_lang == "mermaid":
                    image = render_mermaid(code)
                    if image:
                        add_image(doc, image, f"Figura {figure}. Diagrama de arquitectura / flujo")
                        figure += 1
                    else:
                        add_code_block(doc, code)
                else:
                    add_code_block(doc, code)
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        # Tablas
        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            header = split_row(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            add_table(doc, header, rows)
            continue

        # Imágenes
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            alt, relative = image_match.groups()
            add_image(doc, (base / relative).resolve(), alt or None)
            index += 1
            continue

        # Encabezados
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1:
                if not first_heading:
                    doc.add_page_break()
                paragraph = doc.add_heading(level=1)
                add_inline(paragraph, text)
                first_heading = False
            elif level == 2:
                paragraph = doc.add_heading(level=2)
                add_inline(paragraph, text)
            elif level == 3:
                paragraph = doc.add_heading(level=3)
                add_inline(paragraph, text)
            else:
                paragraph = doc.add_heading(level=4)
                add_inline(paragraph, text)
            index += 1
            continue

        # Separador
        if stripped in ("---", "***", "___"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            p_pr = paragraph._p.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "CBD5E1")
            p_bdr.append(bottom)
            p_pr.append(p_bdr)
            index += 1
            continue

        # Citas
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.space_after = Pt(6)
            shade(paragraph, "EFF6FF")
            add_inline(paragraph, stripped.lstrip("> ").strip(), size=9.5, color=GREY)
            index += 1
            continue

        # Listas
        bullet = re.match(r"^(\s*)-\s+(.*)$", line)
        numbered = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if bullet:
            level = len(bullet.group(1)) // 2
            style = "List Bullet" if level == 0 else "List Bullet 2"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, bullet.group(2))
            index += 1
            continue
        if numbered:
            level = len(numbered.group(1)) // 2
            style = "List Number" if level == 0 else "List Number 2"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, numbered.group(3))
            index += 1
            continue

        # Párrafo normal
        if stripped:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, stripped)
        index += 1

    return figure


# ---------------------------------------------------------------------------
# Portada, ficha y encabezado
# ---------------------------------------------------------------------------
def build_cover(doc: Document) -> None:
    doc.add_paragraph()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("SOC-Lite")
    run.font.size = Pt(38)
    run.bold = True
    run.font.color.rgb = NAVY

    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Triaje automático de alertas SOC")
    run.font.size = Pt(18)
    run.font.color.rgb = BLUE
    run.bold = True

    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Documentación técnica, funcional y de operación")
    run.font.size = Pt(12)
    run.font.color.rgb = GREY

    add_image(doc, DOCS / "screenshots" / "02-dashboard.png", None, width=5.5)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(
        "Caso de uso 1 (crítico) · FastAPI · PostgreSQL · n8n · LangChain/Ollama · React\n"
        "PYME · mediana empresa · organismo público"
    )
    run.font.size = Pt(9.5)
    run.font.color.rgb = GREY

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"Versión 1.0 · {date.today():%d/%m/%Y}")
    run.font.size = Pt(8.5)
    run.font.color.rgb = LIGHT

    doc.add_page_break()


def build_compliance_page(doc: Document) -> None:
    paragraph = doc.add_heading(level=1)
    add_inline(paragraph, "Ficha del entregable")
    intro = doc.add_paragraph()
    add_inline(
        intro,
        "Esta documentación acompaña a una implementación funcional: esquema PostgreSQL, backend "
        "FastAPI, microservicio LangChain, frontend React, workflow n8n, pruebas y capturas de la "
        "aplicación en ejecución. Cada afirmación técnica se contrastó con el código o con una "
        "ejecución real; las limitaciones se declaran al final de esta ficha.",
    )

    headers = ["Requisito", "Implementación", "Evidencia"]
    rows = [
        ["Stack justificado", "FastAPI + SQLAlchemy async + PostgreSQL 16 + React 18/TS + n8n + LangChain/Ollama", "§7 y requirements/package.json"],
        ["PostgreSQL principal", "13 tablas, 5 vistas, 11 enums, JSONB/INET/ARRAY, triggers", "§2, db/schema.sql"],
        ["Dashboard de métricas", "8 endpoints; KPIs, tendencia, severidad, MTTT, SLA, FP y orígenes", "§3.10, §4.7.2"],
        ["Reportes PDF/Word/Excel", "ReportLab, python-docx y openpyxl; 3 tipos × 3 formatos", "§3.11, pruebas"],
        ["Usuarios, roles y perfiles (RBAC)", "5 roles, 26 permisos, matriz editable y validación por endpoint", "§3.5"],
        ["Bitácora de auditoría", "ENTER/EXIT por módulo + eventos; usuario, fecha, hora, acción, IP y hash chain", "§3.6"],
        ["Arquitectura", "Capas, componentes, interacciones, ADR, despliegue y riesgos", "§1"],
        ["Modelo de datos y semillas", "Diccionario, ER, triggers, 728 alertas y 647 triajes", "§2"],
        ["Backend completo", "7 routers, servicios, middleware, 14 pruebas de integración", "§3"],
        ["Frontend completo", "9 páginas, rutas protegidas, cliente con refresh, build verificado", "§4"],
        ["Flujo de nodos", "21 nodos, 10 operativos, JSON importable y SQL validado", "§5 y n8n/"],
        ["Datasets públicos", "25 fuentes en 4 categorías, licencias y protocolo de evaluación", "§6"],
        ["Stack detallado", "Versiones, alternativas, dimensionamiento y coste", "§7"],
    ]
    add_table(doc, headers, rows)

    paragraph = doc.add_heading(level=2)
    add_inline(paragraph, "Verificación ejecutada")
    add_table(
        doc,
        ["Componente", "Verificación", "Resultado"],
        [
            ["PostgreSQL", "schema.sql + seed.sql en PostgreSQL 16", "13 tablas · 728 alertas · 647 triajes"],
            ["Backend", "pytest (integración, BD real)", "14 passed"],
            ["API", "login, RBAC, ingesta, dedup, triaje, feedback, incidente, reportes, auditoría", "OK"],
            ["Reportes", "PDF, DOCX y XLSX generados y descargados", "Archivos válidos"],
            ["Auditoría", "GET /audit/verify-chain", "Cadena íntegra"],
            ["Frontend", "tsc --noEmit + vite build", "0 errores TS · ~867 kB"],
            ["UI", "Chrome headless", "9 capturas"],
            ["n8n", "Import en n8n:latest + SQL contra PostgreSQL", "21 nodos · SQL válido"],
        ],
    )

    paragraph = doc.add_heading(level=2)
    add_inline(paragraph, "Limitaciones declaradas")
    for text in [
        "Ollama no estaba disponible en el entorno de verificación: el triaje respondió con el motor heurístico (status=degraded). El camino LLM está implementado y se activa al detectar Ollama.",
        "El workflow n8n no se ejecutó de punta a punta porque requiere crear la credencial PostgreSQL y las variables de Slack/osTicket. Sí se importó y se validaron sus dos consultas SQL.",
        "Las APIs de threat intelligence requieren claves; sin ellas se marcan como skipped y baja la confianza.",
        "No hay pruebas E2E de navegador automatizadas ni auditoría formal de accesibilidad; se documentan como siguiente paso.",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        add_inline(paragraph, text)

    doc.add_page_break()


def build_toc(doc: Document) -> None:
    paragraph = doc.add_heading(level=1)
    add_inline(paragraph, "Contenido")
    sections = [
        ("1", "Arquitectura del sistema", "01-arquitectura.md"),
        ("2", "Modelo de datos PostgreSQL y datos semilla", "02-modelo-datos.md"),
        ("3", "Desarrollo del backend", "03-backend.md"),
        ("4", "Desarrollo del frontend", "04-frontend.md"),
        ("5", "Flujo de nodos n8n", "05-workflow-n8n.md"),
        ("6", "Datasets públicos recomendados", "06-datasets-publicos.md"),
        ("7", "Stack tecnológico detallado", "07-stack-tecnologico.md"),
        ("A", "Contrato de API v1", "00-contrato-api.md"),
    ]
    for number, title, filename in sections:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"{number}.  ")
        run.bold = True
        run.font.color.rgb = BLUE
        run = paragraph.add_run(title)
        run.font.size = Pt(11)
        run = paragraph.add_run(f"   ({filename})")
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT
    doc.add_page_break()


def config_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in (
        ("Heading 1", 19, NAVY, 0, 10),
        ("Heading 2", 13.5, BLUE, 14, 6),
        ("Heading 3", 11.5, NAVY, 10, 4),
        ("Heading 4", 10.5, GREY, 8, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def config_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "SOC-Lite · Documentación técnica"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Documento de uso interno · ")
    run.font.size = Pt(8)
    run.font.color.rgb = LIGHT
    add_page_number(paragraph)
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT


def main() -> None:
    doc = Document()
    config_page(doc)
    config_styles(doc)
    build_cover(doc)
    build_toc(doc)
    build_compliance_page(doc)

    figure = 1
    order = [
        ("1. Arquitectura del sistema", DOCS / "01-arquitectura.md"),
        ("2. Modelo de datos PostgreSQL y datos semilla", DOCS / "02-modelo-datos.md"),
        ("3. Desarrollo del backend", DOCS / "03-backend.md"),
        ("4. Desarrollo del frontend", DOCS / "04-frontend.md"),
        ("5. Flujo de nodos n8n", DOCS / "05-workflow-n8n.md"),
        ("6. Datasets públicos recomendados", DOCS / "06-datasets-publicos.md"),
        ("7. Stack tecnológico detallado", DOCS / "07-stack-tecnologico.md"),
        ("A. Contrato de API v1", DOCS / "00-contrato-api.md"),
    ]
    for title, path in order:
        print(f"Procesando {path.name}…")
        print(f"  → {title}")
        figure = parse_markdown(doc, path, figure)

    doc.core_properties.title = "SOC-Lite · Documentación técnica y funcional"
    doc.core_properties.author = "Equipo SOC-Lite"
    doc.core_properties.subject = "Triaje automático de alertas SOC"
    doc.core_properties.keywords = "SOC, SOAR, n8n, LangChain, FastAPI, PostgreSQL, ISO 27001, NIST CSF"
    doc.save(OUT)
    print(f"\nOK → {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    sys.exit(main())
